from __future__ import annotations

from io import BytesIO
from pathlib import PurePosixPath
from typing import Any

from openpyxl import load_workbook

from ingestion.models import ParsedDocument


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""

    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except TypeError:
            pass

    return str(value).strip()


def _normalize_headers(values: tuple[Any, ...]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}

    for index, value in enumerate(values, start=1):
        header = _cell_to_text(value) or f"column_{index}"

        if header in seen:
            seen[header] += 1
            header = f"{header}_{seen[header]}"
        else:
            seen[header] = 1

        headers.append(header)

    return headers


def parse_xlsx_bytes(
    data: bytes,
    source_path: str,
) -> list[ParsedDocument]:
    workbook = load_workbook(
        filename=BytesIO(data),
        read_only=True,
        data_only=True,
    )

    source_file = PurePosixPath(source_path).name
    domain = source_path.split("/", 1)[0] if "/" in source_path else ""

    documents: list[ParsedDocument] = []

    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)

        header_row = None
        header_row_number = None

        for row_number, row in enumerate(rows, start=1):
            if any(value is not None and str(value).strip() for value in row):
                header_row = row
                header_row_number = row_number
                break

        if header_row is None or header_row_number is None:
            continue

        headers = _normalize_headers(header_row)

        for row_number, row in enumerate(
            rows,
            start=header_row_number + 1,
        ):
            if not any(
                value is not None and str(value).strip()
                for value in row
            ):
                continue

            values = list(row)

            if len(values) < len(headers):
                values.extend([None] * (len(headers) - len(values)))

            record = {
                header: _cell_to_text(value)
                for header, value in zip(headers, values)
            }

            content = "\n".join(
                f"{key}: {value}"
                for key, value in record.items()
                if key
            )

            safe_sheet_name = worksheet.title.replace("#", "_")

            documents.append(
                ParsedDocument(
                    document_id=(
                        f"{source_path}"
                        f"#sheet={safe_sheet_name}"
                        f"#row={row_number}"
                    ),
                    source_path=source_path,
                    file_type="xlsx",
                    content=content,
                    metadata={
                        "domain": domain,
                        "source_file": source_file,
                        "record_type": "xlsx_row",
                        "sheet_name": worksheet.title,
                        "row_number": str(row_number),
                    },
                )
            )

    workbook.close()

    return documents
